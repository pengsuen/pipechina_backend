"""Build mixed-layout documents with image-only scans and rule-grounded questions.

Use the bundled artifact Python runtime. Output must be a new directory.
"""

# ruff: noqa: E501 -- document prose is kept as complete paragraphs.
# mypy: disable-error-code="import-untyped,import-not-found"

from __future__ import annotations

import argparse
import base64
import hashlib
import html
import io
import json
import os
import random
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

from knowledge_corpus_content import catalog, content, records

RUNTIME = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
FONT = os.environ.get("KNOWLEDGE_CORPUS_FONT", "/System/Library/Fonts/Supplemental/Songti.ttc")
MIMES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "html": "text/html",
    "md": "text/markdown",
    "txt": "text/plain",
}


def format_for(n):
    # Exactly 60 primary documents: PDF 36, DOCX 12, XLSX 6, HTML 3, MD 2, TXT 1.
    if n % 10 <= 5:
        return "pdf", ["illustrated", "scan", "two_column", "mixed", "landscape", "two_column"][
            n % 10
        ]
    if n % 10 in (6, 7):
        return "docx", "word_tables"
    if n % 10 == 8:
        return "xlsx", "workbook"
    return [
        ("html", "web_article"),
        ("md", "structured_text"),
        ("html", "web_article"),
        ("txt", "plain_text"),
        ("md", "structured_text"),
        ("html", "web_article"),
    ][n // 10]


def image_assets(spec, folder):
    """Draw labeled technical diagrams and a real raster chart, not decorative photos."""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(FONT, 24, index=6)
    small = ImageFont.truetype(FONT, 20, index=6)
    im = Image.new("RGB", (1280, 480), "white")
    d = ImageDraw.Draw(im)
    d.text((35, 20), spec["station"] + "资料与证据流转图", fill="#173b52", font=font)
    labels = ["现场观察", "原始记录", "专业复核", "验收归档"]
    for i, label in enumerate(labels):
        x = 35 + i * 310
        d.rounded_rectangle(
            (x, 130, x + 230, 225), radius=8, fill="#edf3f7", outline="#426880", width=3
        )
        d.text((x + 62, 160), label, fill="#173b52", font=font)
        if i < 3:
            d.line((x + 235, 175, x + 300, 175), fill="#426880", width=4)
            d.polygon([(x + 300, 175), (x + 285, 165), (x + 285, 185)], fill="#426880")
        d.text(
            (x, 255),
            ["图像 时间 位号", "编号 摘要 来源", "条件 反证 例外", "版本 权限 目录"][i],
            fill="#333333",
            font=small,
        )
    d.line((760, 330, 150, 330, 150, 235), fill="#a96526", width=3)
    d.text((310, 350), "补证退回 保留原编号和首次接收时间", fill="#80501e", font=small)
    diagram = folder / f"{spec['code']}-flow.png"
    im.save(diagram)
    chart = Image.new("RGB", (1280, 480), "white")
    d = ImageDraw.Draw(chart)
    d.text((35, 16), "记录复核量与待补正量 月度趋势", fill="#173b52", font=font)
    d.line((90, 85, 90, 385, 1220, 385), fill="#555555", width=2)
    values = []
    for i in range(6):
        total = 35 + ((i * 7 + spec["ordinal"] * 3) % 24)
        pending = 3 + (spec["ordinal"] + i * 3) % 9
        values.append([f"{i + 1}月", total, pending])
        x = 140 + i * 180
        d.rectangle((x, 385 - total * 4, x + 52, 385), fill="#335e79")
        d.rectangle((x + 58, 385 - pending * 4, x + 110, 385), fill="#bd7f39")
        d.text((x, 395), f"{i + 1}月", font=small, fill="#333333")
        d.text((x, 350 - total * 4), str(total), font=small, fill="#333333")
        d.text((x + 65, 350 - pending * 4), str(pending), font=small, fill="#80501e")
    d.text((820, 18), "蓝 复核量   棕 待补正量", fill="#333333", font=small)
    trend = folder / f"{spec['code']}-trend.png"
    chart.save(trend)
    return diagram, trend, values


def pdf_file(path, spec, version, sections, layout, assets, scratch):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        Image,
        NextPageTemplate,
        PageBreak,
        PageTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    pdfmetrics.registerFont(TTFont("CorpusSong", FONT, subfontIndex=6))
    wide = layout == "landscape"
    size = landscape(A4) if wide else A4
    w, h = size
    usable = w - 96
    body = ParagraphStyle(
        "body", fontName="CorpusSong", fontSize=10.5, leading=17, spaceAfter=8, wordWrap="CJK"
    )
    head = ParagraphStyle(
        "heading",
        parent=body,
        fontSize=14,
        leading=21,
        spaceBefore=12,
        spaceAfter=8,
        keepWithNext=True,
    )
    title = ParagraphStyle("title", parent=head, fontSize=21, leading=31, spaceBefore=0)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=12, spaceAfter=0)

    def p(text, style=body):
        return Paragraph(html.escape(str(text)), style)

    def footer(canvas, doc):
        canvas.setFont("CorpusSong", 8)
        canvas.setFillColor(colors.HexColor("#5a6268"))
        canvas.drawString(48, h - 26, f"青衡管网运营有限公司  {spec['code']}  V{version}.0")
        canvas.drawString(48, 26, "运行管理部  受控文件")
        canvas.drawRightString(w - 48, 26, f"第 {doc.page} 页")

    single = Frame(48, 46, usable, h - 94, id="single", leftPadding=0, rightPadding=0)
    gap = 23
    col = (usable - gap) / 2
    frames = [
        Frame(48, 46, col, h - 94, id="left", leftPadding=0, rightPadding=0),
        Frame(48 + col + gap, 46, col, h - 94, id="right", leftPadding=0, rightPadding=0),
    ]
    doc = BaseDocTemplate(
        str(path), pagesize=size, title=spec["title"], author="青衡管网运行管理部"
    )
    doc.addPageTemplates(
        [
            PageTemplate(id="single", frames=[single], onPage=footer),
            PageTemplate(id="columns", frames=frames, onPage=footer),
        ]
    )

    def table(rows, widths=None):
        t = Table(
            [[p(c, cell) for c in row] for row in rows],
            colWidths=widths,
            repeatRows=1,
            hAlign="LEFT",
        )
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dce6ed")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#f5f7f9")],
                    ),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d9d9d9")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        return t

    story = [
        p(spec["title"], title),
        p(f"文件编号 {spec['code']}    版本 {version}.0"),
        p(
            f"生效日期 {'2025年1月1日' if version == 1 else '2026年7月1日'}    编制 {spec['owner']}    审核 林致远"
        ),
        p("文件摘要", head),
        p(sections[0]["paragraphs"][0]),
        p("修订记录", head),
        table(
            [
                ["版本", "生效日期", "主要变化"],
                ["1.0", "2025-01-01", "初审3个工作日；保存5年；月度抽查10%"],
                *(
                    [["2.0", "2026-07-01", "初审2个工作日；保存8年；月度抽查15%；新增存量衔接条款"]]
                    if version == 2
                    else []
                ),
            ],
            [usable * 0.12, usable * 0.22, usable * 0.66],
        ),
        p("内容索引", head),
        *(
            [p("；".join(s["heading"] for s in sections))]
            if wide
            else [p(s["heading"]) for s in sections]
        ),
        p("附件A 规则检查表  附件B 记录台账  附件C 图像与数据"),
    ]
    if layout == "two_column":
        story.append(NextPageTemplate("columns"))
    story.append(PageBreak())
    for s in sections:
        story.append(p(s["heading"], head))
        story.extend(p(text) for text in s["paragraphs"])
    if layout == "two_column":
        story.append(NextPageTemplate("single"))
    story += [PageBreak(), p("附件A 规则检查表", head)]
    rules = [["检查项", "通过条件", "不满足时处理"]]
    for i, text in enumerate(sections[2]["paragraphs"]):
        rules.append(
            [f"A{i + 1:02d}", text.split("。", 1)[0] + "。", "补充原始记录并由独立人员复核"]
        )
    story.append(table(rules, [usable * 0.1, usable * 0.68, usable * 0.22]))
    story += [
        PageBreak(),
        p("附件B 记录台账", head),
        p("统计范围为本次复核资料；同一事项只按主编号计数。"),
        table(records(spec), [usable * x for x in [0.06, 0.16, 0.23, 0.12, 0.10, 0.10, 0.23]]),
    ]
    story += [
        PageBreak(),
        p("附件C 图像与数据", head),
        Image(str(assets[0]), width=usable, height=usable * 480 / 1280),
        p("图1 资料流转与补证退回关系。退回时保留原编号，首次接收时间不重置。"),
        Spacer(1, 12),
        Image(str(assets[1]), width=usable, height=usable * 480 / 1280),
        p("图2 月度记录复核与待补正情况。待补正量包含在复核量内，不重复累加。"),
        table([["月份", "复核量", "待补正量"], *assets[2]], [usable / 3] * 3),
    ]
    doc.build(story, canvasmaker=lambda *a, **kw: Canvas(*a, **{**kw, "invariant": 1}))
    if layout in {"scan", "mixed"}:
        rasterize_pdf(path, scratch, mixed=layout == "mixed", seed=spec["ordinal"])


def rasterize_pdf(path, scratch, mixed, seed):
    from PIL import Image, ImageEnhance, ImageFilter
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    source = PdfReader(path)
    writer = PdfWriter()
    rng = random.Random(seed)
    for i, page in enumerate(source.pages):
        if mixed and i % 2 == 0:
            writer.add_page(page)
            continue
        prefix = scratch / f"{path.stem}-{i}"
        subprocess.run(
            [
                str(RUNTIME / "bin/override/pdftoppm"),
                "-f",
                str(i + 1),
                "-l",
                str(i + 1),
                "-singlefile",
                "-r",
                "135",
                "-png",
                str(path),
                str(prefix),
            ],
            check=True,
            capture_output=True,
        )
        im = Image.open(prefix.with_suffix(".png")).convert("L")
        im = ImageEnhance.Contrast(im).enhance(0.88)
        im = im.rotate(
            rng.choice([-0.35, 0.25, 0.45]), resample=Image.Resampling.BICUBIC, fillcolor=248
        )
        im = im.filter(ImageFilter.GaussianBlur(0.18)).convert("RGB")
        buf = io.BytesIO()
        c = Canvas(
            buf, pagesize=(float(page.mediabox.width), float(page.mediabox.height)), invariant=1
        )
        c.drawImage(
            ImageReader(im),
            0,
            0,
            width=float(page.mediabox.width),
            height=float(page.mediabox.height),
        )
        c.save()
        buf.seek(0)
        writer.add_page(PdfReader(buf).pages[0])
        prefix.with_suffix(".png").unlink()
    result = io.BytesIO()
    writer.write(result)
    path.write_bytes(result.getvalue())


def word_file(path, spec, version, sections, assets):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.add_paragraph(spec["title"], "Title")
    doc.add_paragraph(f"{spec['code']}  版本 {version}.0  编制 {spec['owner']}  审核 林致远")
    doc.add_paragraph("生效日期 " + ("2025-01-01" if version == 1 else "2026-07-01"))
    sec.header.paragraphs[0].text = "青衡管网运营有限公司  运行管理部"
    footer = sec.footer.paragraphs[0]
    footer.add_run("受控文件  第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    for s in sections:
        doc.add_heading(s["heading"], 1)
        for text in s["paragraphs"]:
            doc.add_paragraph(text)
    doc.add_page_break()
    doc.add_heading("附件A 复核台账", 1)
    rows = records(spec, 36)
    t = doc.add_table(rows=1, cols=7)
    t.autofit = False
    widths = [0.8, 2.6, 3.2, 1.8, 1.4, 1.4, 5.4]
    for col, width in zip(t.columns, widths, strict=True):
        col.width = Cm(width)
    for j, text in enumerate(rows[0]):
        t.rows[0].cells[j].text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows[1:]:
        for c, text in zip(t.add_row().cells, row, strict=True):
            c.text = text
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCE6ED" if i == 0 else "FFFFFF" if i % 2 else "F5F7F9")
            c._tc.get_or_add_tcPr().append(shade)
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    doc.add_page_break()
    doc.add_heading("附件B 图像与资料流转", 1)
    doc.add_picture(str(assets[0]), width=Cm(16))
    doc.add_paragraph("图1 资料流转与补证关系", style="Caption")
    doc.add_picture(str(assets[1]), width=Cm(16))
    doc.add_paragraph("图2 记录质量月度趋势", style="Caption")
    doc.save(path)


def build(output):
    if output.exists() and any(output.iterdir()):
        raise ValueError("output must be empty; no existing corpus is overwritten")
    output.mkdir(parents=True, exist_ok=True)
    docs, questions, jobs = [], [], []
    with tempfile.TemporaryDirectory(prefix="qh-rich-") as tmp:
        scratch = Path(tmp)
        for spec in catalog():
            suffix, layout = format_for(spec["ordinal"])
            folder = output / (spec["kind"] + "s")
            folder.mkdir(exist_ok=True)
            assets = image_assets(spec, scratch)
            entries = []
            for version in range(1, spec["versions"] + 1):
                path = folder / f"{spec['code']}-v{version}.{suffix}"
                sections = content(spec, version)
                if suffix == "docx":
                    sections[-1]["paragraphs"][-1] = (
                        "附件A为复核台账，附件B包含资料流转图与月度趋势图。台账编号应与正文引用一致，复核结果以实际检查范围为准。"
                    )
                elif suffix == "xlsx":
                    sections[-1]["paragraphs"][-1] = (
                        "本工作簿包含复核汇总、明细台账、规则与例外三个工作表。汇总数由明细状态计算，月度趋势图引用汇总页数据。规则与例外工作表保存完整执行口径。"
                    )
                elif suffix in {"md", "txt"}:
                    sections[-1]["paragraphs"][-1] = (
                        "引用本记录时注明文件编号、版本和章节。相关原始证据由资料管理员按事项编号归集，尚未取得的附件不得填报为已核对。"
                    )
                chars = sum(len(p) for s in sections for p in s["paragraphs"])
                if suffix == "pdf":
                    pdf_file(path, spec, version, sections, layout, assets, scratch)
                elif suffix == "docx":
                    word_file(path, spec, version, sections, assets)
                elif suffix == "xlsx":
                    jobs.append(
                        dict(
                            path=str(path.resolve()),
                            spec=spec,
                            version=version,
                            sections=sections,
                            rows=records(spec, 72),
                            trend=assets[2],
                        )
                    )
                else:
                    text = (
                        f"# {spec['title']}\n\n{spec['code']}  版本 {version}.0\n\n"
                        + "\n\n".join(
                            "## " + s["heading"] + "\n\n" + "\n\n".join(s["paragraphs"])
                            for s in sections
                        )
                    )
                    if suffix == "html":
                        text = (
                            "<!doctype html><html lang='zh'><meta charset='utf-8'><title>"
                            + html.escape(spec["title"])
                            + "</title><style>body{max-width:1000px;margin:45px auto;font:17px/1.8 serif}article{columns:2;column-gap:40px}h2{break-after:avoid}img{width:100%}table{border-collapse:collapse}td,th{border:1px solid #aaa;padding:8px}</style><h1>"
                            + spec["title"]
                            + f"</h1><p>{spec['code']}  版本 {version}.0</p><article>"
                            + "".join(
                                "<h2>"
                                + s["heading"]
                                + "</h2>"
                                + "".join("<p>" + html.escape(p) + "</p>" for p in s["paragraphs"])
                                for s in sections
                            )
                            + "</article><table>"
                            + "".join(
                                "<tr>"
                                + "".join("<td>" + html.escape(c) + "</td>" for c in row)
                                + "</tr>"
                                for row in records(spec, 24)
                            )
                            + "</table>"
                            + "".join(
                                '<figure><img alt="'
                                + caption
                                + '" src="data:image/png;base64,'
                                + base64.b64encode(asset.read_bytes()).decode()
                                + '"><figcaption>'
                                + caption
                                + "</figcaption></figure>"
                                for asset, caption in [
                                    (assets[0], "图1 资料流转关系"),
                                    (assets[1], "图2 月度记录趋势"),
                                ]
                            )
                            + "</html>"
                        )
                    path.write_text(text, encoding="utf-8")
                entries.append(
                    dict(
                        number=version,
                        filename=path.name,
                        path=str(path.relative_to(output)),
                        mime_type=MIMES[suffix],
                        effective_from="2025-01-01T00:00:00+08:00"
                        if version == 1
                        else "2026-07-01T00:00:00+08:00",
                        equipment_models=[f"{spec['station_code']}-{spec['index']:02d}"],
                        authority="内部",
                        layout=layout,
                        body_characters=chars,
                    )
                )
                for name, expected in [
                    ("一般事项初审时限", f"{3 if version == 1 else 2}个工作日"),
                    ("完整证据保存期限", f"{5 if version == 1 else 8}年"),
                    ("月度分层抽查比例", f"{10 if version == 1 else 15}%"),
                ]:
                    questions.append(
                        dict(
                            question=f"{spec['code']} 第{version}版的{name}是多少？",
                            document_code=spec["code"],
                            version=version,
                            expected_answer=expected,
                            expected_quote=next(
                                p for p in sections[0]["paragraphs"] if "本版一般事项" in p
                            ),
                            should_refuse=False,
                            split="held_out" if spec["ordinal"] % 4 == 0 else "development",
                            locator_hint="一 文件适用范围与管理边界",
                        )
                    )
            questions.append(
                dict(
                    question=f"{spec['code']} 是否允许通过撤回重提清除逾期记录？",
                    document_code=spec["code"],
                    version=spec["versions"],
                    expected_answer="不允许",
                    expected_quote="不允许通过撤回重提清除逾期记录。",
                    should_refuse=False,
                    split="development",
                )
            )
            questions.append(
                dict(
                    question=f"{spec['code']} 中设备未披露的制造商序列号是什么？",
                    document_code=spec["code"],
                    version=spec["versions"],
                    should_refuse=True,
                    reason="文件未提供制造商序列号",
                    split="held_out",
                )
            )
            docs.append(
                dict(code=spec["code"], title=spec["title"], kind=spec["kind"], versions=entries)
            )
            print(spec["code"], suffix, layout, flush=True)
        if jobs:
            jobs_path = scratch / "workbooks.json"
            jobs_path.write_text(json.dumps(jobs, ensure_ascii=False), encoding="utf-8")
            # Resolve the documented artifact package from the bundled dependency root.
            runner = scratch / "build_workbooks.mjs"
            shutil.copyfile(Path(__file__).with_name("build_corpus_workbooks.mjs"), runner)
            (scratch / "node_modules").symlink_to(
                RUNTIME / "node/node_modules", target_is_directory=True
            )
            subprocess.run(
                [
                    str(RUNTIME / "node/bin/node"),
                    str(runner),
                    str(jobs_path),
                    str(output.resolve()),
                ],
                check=True,
            )
    from pypdf import PdfReader

    for d in docs:
        for v in d["versions"]:
            path = output / v["path"]
            data = path.read_bytes()
            v.update(size_bytes=len(data), sha256=hashlib.sha256(data).hexdigest())
            if path.suffix == ".pdf":
                reader = PdfReader(path)
                v["pages"] = len(reader.pages)
                v["image_only_pages"] = sum(
                    not (p.extract_text() or "").strip() for p in reader.pages
                )
    allv = [v for d in docs for v in d["versions"]]
    manifest = dict(
        schema_version=3,
        corpus_name="青衡管网生产运行知识库",
        document_count=len(docs),
        version_count=len(allv),
        formats=dict(Counter(Path(v["path"]).suffix for v in allv)),
        layouts=dict(Counter(v["layout"] for v in allv)),
        documents=docs,
        questions=questions,
    )
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    lines = [
        "# 知识库语料目录",
        "",
        f"60 个主题，75 个版本文件，{len(questions)} 条有依据的验收问题。",
        "",
        "|文号|主题|格式与结构|版本|",
        "|---|---|---|---|",
    ]
    for d in docs:
        v = d["versions"][-1]
        lines.append(
            f"|{d['code']}|[{d['title']}]({v['path']})|{Path(v['path']).suffix} / {v['layout']}|{len(d['versions'])}|"
        )
    lines += [
        "",
        "扫描PDF的扫描页只有图像，没有隐藏文字层。混合PDF交替包含文字页与扫描页。",
        "",
        "批量导入使用 scripts/import_knowledge_corpus.py。manifest.json 只登记业务文件，目录及校验报告不上传。",
        "",
        "生成命令：使用包含 reportlab、python-docx、Pillow、pypdf 的运行环境执行 scripts/build_rich_knowledge_corpus.py --output 新目录。XLSX 使用 Codex bundled artifact-tool，扫描使用 pdftoppm。",
    ]
    (output / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = build(args.output)
    print(
        json.dumps(
            {k: result[k] for k in ["document_count", "version_count", "formats", "layouts"]},
            ensure_ascii=False,
        )
    )
